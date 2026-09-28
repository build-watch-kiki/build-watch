from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from buildwatch.infrastructure.database.models.core.projects import (
    ProjectTypesORM,
    ProjectTypesWorkTypesORM,
    WorkTypesORM,
)


def _normalize_code(value) -> str | None:
    """Нормализация кода работы из ячейки Excel"""

    if value is None:
        return None
    if isinstance(value, datetime):
        # Excel хранит 10.01. как 2025-01-10 (day=10, month=01)
        return f"10.{value.month:02d}."
    s = str(value).strip()
    return s if s else None


def _parent_code(code: str, existing_codes: set[str]) -> str | None:
    """Находит родителя по префиксу кода. '10.11.1.' -> '10.11.', '10.01.' -> '10.'"""
    if not code:
        return None
    stripped = code.rstrip(".")
    if "." not in stripped:
        return None
    # пытаемся отбросить последний сегмент
    parts = stripped.split(".")
    for i in range(len(parts) - 1, 0, -1):
        cand_dot = ".".join(parts[:i]) + "."
        if cand_dot in existing_codes:
            return cand_dot
        cand = ".".join(parts[:i])
        if cand in existing_codes:
            return cand
    # fallback на первый сегмент
    if parts[0] in existing_codes:
        return parts[0]
    if parts[0] + "." in existing_codes:
        return parts[0] + "."
    return None


async def add_work_types(
    xlsx_path: Path,
    session: AsyncSession,
) -> list[WorkTypesORM]:
    """Заполнение справочника типов работ и связей из Excel"""

    if not xlsx_path.exists():
        raise FileNotFoundError(f"Excel файл не найден: {xlsx_path}")

    # мапа project_type name -> id для линковки
    pt_rows = (await session.execute(select(ProjectTypesORM))).scalars().all()
    pt_map = {pt.name: pt.id for pt in pt_rows}
    if not pt_map:
        raise ValueError(
            "Справочник project_types пуст — сначала вызови add_project_types"
        )

    # порядок колонок C-K должен совпадать с заголовком
    header_order = [
        "Жильё",
        "Образование",
        "Здравоохранение",
        "Спорт",
        "Культура",
        "Административные здания",
        "ДОУ",
        "Офисно-деловой центр",
        "Дороги",
    ]

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        if ws is None:
            raise ValueError("Нет активного листа")

        # читаем все строки 4..380: (code, name, flags[9])
        rows = []
        for row in ws.iter_rows(
            min_row=4, max_row=380, min_col=1, max_col=11, values_only=True
        ):
            raw_code, name = row[0], row[1]
            flags = row[2:11]  # C-K
            code = _normalize_code(raw_code)
            if name is None or not str(name).strip():
                continue
            rows.append((code, str(name).strip(), flags))
    finally:
        wb.close()

    # идемпотентность по (code, name) — code может быть None, тогда по name
    existing = (await session.execute(select(WorkTypesORM))).scalars().all()
    existing_keys = set((wt.code, wt.name) for wt in existing)
    code_to_id: dict[str, int] = {}
    for wt in existing:
        if wt.code:
            code_to_id.setdefault(wt.code, wt.id)

    last_code_id: int | None = None
    created: list[WorkTypesORM] = []
    # для линковки копим пары (work_type_id, project_type_id)
    links_to_create: list[tuple[int, int]] = []

    # существующие линки для идемпотентности
    existing_links = set(
        (
            await session.execute(
                select(
                    ProjectTypesWorkTypesORM.project_type_id,
                    ProjectTypesWorkTypesORM.work_type_id,
                )
            )
        ).all()
    )

    for code, name, flags in rows:
        key = (code, name)
        if key in existing_keys:
            # уже есть — находим его id для линковки и для parent цепочки
            wt = next(w for w in existing if (w.code, w.name) == key)
            wt_id = wt.id
            if code:
                code_to_id[code] = wt_id
                last_code_id = wt_id
            # линки всё равно проверим ниже
        else:
            # определяем parent_id
            if code is None:
                parent_id = last_code_id
            else:
                pcode = _parent_code(code, set(code_to_id.keys()))
                parent_id = code_to_id.get(pcode) if pcode else None
            wt = WorkTypesORM(name=name, code=code, parent_id=parent_id)
            session.add(wt)
            await session.flush()  # получаем id
            wt_id = wt.id
            created.append(wt)
            existing_keys.add(key)
            if code:
                code_to_id[code] = wt_id
                last_code_id = wt_id
            existing.append(wt)

        # линки project_types_work_types по матрице
        for idx, flag in enumerate(flags):
            if flag is None:
                continue
            if str(flag).strip() != "˅":
                continue
            pt_name = header_order[idx]
            pt_id = pt_map.get(pt_name)
            if pt_id is None:
                continue
            if (pt_id, wt_id) in existing_links:
                continue
            links_to_create.append((pt_id, wt_id))
            existing_links.add((pt_id, wt_id))

    if links_to_create:
        link_objs = [
            ProjectTypesWorkTypesORM(project_type_id=pt, work_type_id=wt)
            for pt, wt in links_to_create
        ]
        session.add_all(link_objs)
        await session.flush()
        print(f"Добавлено связей project_types_work_types: {len(link_objs)}")
    else:
        print("Новых связей project_types_work_types нет")

    if created:
        print(f"Добавлено work_types: {len(created)}")
    else:
        print("Новых work_types нет (все 377 уже в БД)")

    return created

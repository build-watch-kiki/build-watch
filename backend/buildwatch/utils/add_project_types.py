import logging
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from buildwatch.infrastructure.database.models.core.projects import ProjectTypesORM

logger = logging.getLogger(__name__)


def _read_project_types(xlsx_path: Path) -> list[str]:
    """Чтение типов проектов из строки C3:K3 Excel-листа"""

    if not xlsx_path.exists():
        raise FileNotFoundError(f"Excel файл не найден: {xlsx_path}")

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        if ws is None:
            raise ValueError("В Excel файле нет активного листа")
        row = next(
            ws.iter_rows(min_row=3, max_row=3, min_col=3, max_col=11, values_only=True)
        )
        names = [str(v).strip() for v in row if v is not None and str(v).strip()]
        if not names:
            raise ValueError(f"Не найдены типы проектов в {xlsx_path} (C3:K3 пусто)")
        return names
    finally:
        wb.close()


async def _insert_project_types(
    session: AsyncSession,
    names: list[str],
) -> list[ProjectTypesORM]:
    """Идемпотентная вставка типов проектов"""

    existing = set(
        (await session.execute(select(ProjectTypesORM.name))).scalars().all()
    )
    to_insert = [n for n in names if n not in existing]
    if not to_insert:
        print(f"Все типы проектов уже в БД ({len(existing)} шт)")
        return []

    objects = [ProjectTypesORM(name=name) for name in to_insert]
    session.add_all(objects)
    await session.flush()
    print(f"Добавлено {len(objects)} типов проектов: {to_insert}")
    return objects


async def add_project_types(
    xlsx_path: Path,
    session: AsyncSession,
) -> list[ProjectTypesORM]:
    """Заполняет справочник project_types из Excel (C3:K3)"""
    names = _read_project_types(xlsx_path)
    return await _insert_project_types(session, names)
